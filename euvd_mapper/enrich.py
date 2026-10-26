"""Add EUVD results to a BOM workbook without rewriting the source sheets."""

from __future__ import annotations

import sys
from collections import Counter
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from euvd_mapper import __version__
from euvd_mapper.aliases import normalize_name
from euvd_mapper.http import CacheMiss
from euvd_mapper.match import (
    ComponentResult,
    format_score,
    match_component,
    rollup_basis,
    rollup_verdict,
)

PURL_HEADERS = ("package url", "packageurl", "purl", "package_url")
ORIGIN_HEADERS = ("origin id", "origin name/id", "origin name / id")
NAME_HEADERS = ("component name", "product name", "product", "name")
VERSION_HEADERS = ("component version name", "component version", "version name", "version")
RESULT_HEADERS = (
    "EUVD IDs",
    "CVE/GHSA aliases",
    "Max severity/score",
    "Version verdict",
    "Match basis",
    "Evidence/notes",
)
DETAIL_HEADERS = (
    "Component name",
    "Version",
    "PURL",
    "EUVD ID",
    "Aliases",
    "Severity",
    "Score",
    "Version verdict",
    "Match basis",
    "Cross check",
    "Evidence",
    "Source row",
)
EUVD_LINK = "https://euvd.enisa.europa.eu/vulnerability/"
HEADER_SCAN = 40


class WorkbookError(ValueError):
    pass


def enrich_workbook(
    input_path: str | Path,
    output_path: str | Path,
    http,
    sheet: str | None = None,
    name_column: str | None = None,
    version_column: str | None = None,
    purl_column: str | None = None,
    matcher=match_component,
    now=None,
) -> list[str]:
    source = Path(input_path)
    target = Path(output_path)
    if source.resolve() == target.resolve():
        raise WorkbookError("refusing to overwrite the input workbook")
    if not any((sheet, name_column, version_column, purl_column)) and _is_project_report(source):
        return _enrich_project_report(source, target, http, now)
    workbook = load_workbook(source)
    explicit = tuple(item for item in (name_column, version_column, purl_column) if item)
    worksheet = _choose_sheet(workbook, sheet, explicit)
    header_row, headers = _header_row(worksheet, explicit)
    columns = _columns(headers, name_column, version_column, purl_column)
    if columns["name"] is None and columns["purl"] is None:
        raise WorkbookError(f"no component name or purl column on {worksheet.title}")

    start = _result_start(headers)
    _write_header(worksheet, header_row, start, headers)
    rows = _read_rows(worksheet, header_row, columns)
    resolved, unresolved = _lookup_rows(rows, http, matcher)
    for row_number, result in resolved.items():
        _write_result(worksheet, row_number, start, result)
    _write_detail(workbook, resolved)
    _write_run(workbook, resolved, unresolved, source, http)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return _summary(list(resolved.values()), unresolved, target)


def _is_project_report(source: Path) -> bool:
    from euvd_mapper.bdreport import is_blackduck_project_report

    workbook = load_workbook(source, data_only=False)
    try:
        return is_blackduck_project_report(workbook)
    finally:
        workbook.close()


def _enrich_project_report(source: Path, target: Path, http, now) -> list[str]:
    from euvd_mapper.bdreport import (
        EUVD_HEADERS,
        EUVD_SHEET,
        SUMMARY_SHEET,
        VULN_SHEET,
        build_report,
        read_bom,
        read_vulnerabilities,
    )
    from euvd_mapper.xlsxinsert import insert_euvd_sheets

    workbook = load_workbook(source)
    try:
        report = build_report(http, read_vulnerabilities(workbook), read_bom(workbook), now=now)
    finally:
        workbook.close()
    insert_euvd_sheets(
        source,
        target,
        report.banner,
        list(EUVD_HEADERS),
        report.euvd_rows,
        report.summary_rows,
        euvd_sheet=EUVD_SHEET,
        summary_sheet=SUMMARY_SHEET,
        after_sheet=VULN_SHEET,
    )
    return report.stdout + [f"wrote {target}"]


def _choose_sheet(workbook, sheet: str | None, explicit: tuple[str, ...] = ()):
    if sheet:
        if sheet not in workbook.sheetnames:
            raise WorkbookError(f"sheet not found: {sheet}")
        return workbook[sheet]
    ranked = []
    for worksheet in workbook.worksheets:
        try:
            _row, headers = _header_row(worksheet, explicit)
        except WorkbookError:
            continue
        score = _header_score(headers)
        if any(word in worksheet.title.casefold() for word in ("bom", "component")):
            score += 0.5
        ranked.append((score, worksheet))
    if not ranked:
        raise WorkbookError("could not find a header row with component columns")
    ranked.sort(key=lambda item: item[0], reverse=True)
    return ranked[0][1]


def _header_row(worksheet, explicit: tuple[str, ...] = ()) -> tuple[int, dict[int, str]]:
    best: tuple[float, int, dict[int, str]] | None = None
    last = min(HEADER_SCAN, worksheet.max_row or 1)
    for row in range(1, last + 1):
        headers = {}
        for col in range(1, (worksheet.max_column or 1) + 1):
            value = worksheet.cell(row, col).value
            if isinstance(value, str) and value.strip():
                headers[col] = value.strip()
        score = _header_score(headers)
        names = {_norm_header(value) for value in headers.values()}
        if names & {_norm_header(item) for item in explicit}:
            score += 2
        if best is None or score > best[0]:
            best = (score, row, headers)
    if best is None or best[0] < 2:
        raise WorkbookError(f"no component header on {worksheet.title}")
    return best[1], best[2]


def _header_score(headers: dict[int, str]) -> float:
    names = {_norm_header(value) for value in headers.values()}
    score = 0
    if names & set(NAME_HEADERS):
        score += 2
    if names & set(PURL_HEADERS) or names & set(ORIGIN_HEADERS):
        score += 2
    if names & set(VERSION_HEADERS):
        score += 1
    return score


def _columns(headers, name_column, version_column, purl_column) -> dict[str, int | None]:
    return {
        "name": _pick(headers, [name_column] if name_column else NAME_HEADERS),
        "version": _pick(headers, [version_column] if version_column else VERSION_HEADERS),
        "purl": _pick(headers, [purl_column] if purl_column else PURL_HEADERS),
        "origin": _pick(headers, ORIGIN_HEADERS),
    }


def _pick(headers: dict[int, str], wanted) -> int | None:
    wanted_norm = [_norm_header(item) for item in wanted if item]
    by_name = {_norm_header(value): col for col, value in headers.items()}
    for name in wanted_norm:
        if name in by_name:
            return by_name[name]
    return None


def _norm_header(value: str) -> str:
    return " ".join(value.strip().casefold().replace("_", " ").split())


def _result_start(headers: dict[int, str]) -> int:
    for col, value in headers.items():
        if _norm_header(value) == "euvd ids":
            return col
    return max(headers) + 1


def _write_header(worksheet, header_row: int, start: int, headers: dict[int, str]) -> None:
    source = worksheet.cell(header_row, max(headers))
    for offset, title in enumerate(RESULT_HEADERS):
        cell = worksheet.cell(header_row, start + offset, title)
        if source.has_style:
            cell.font = copy(source.font)
            cell.fill = copy(source.fill)
            cell.border = copy(source.border)
            cell.alignment = copy(source.alignment)


def _read_rows(worksheet, header_row: int, columns: dict[str, int | None]) -> list[dict]:
    rows = []
    for row in range(header_row + 1, (worksheet.max_row or header_row) + 1):
        name = _text(worksheet, row, columns["name"])
        version = _text(worksheet, row, columns["version"])
        purl = _as_purl(_text(worksheet, row, columns["purl"]))
        if purl is None:
            purl = _as_purl(_text(worksheet, row, columns["origin"]))
        if not name and not version and not purl:
            continue
        rows.append({"row": row, "name": name, "version": version, "purl": purl})
    return rows


def _text(worksheet, row: int, col: int | None) -> str:
    if col is None:
        return ""
    value = worksheet.cell(row, col).value
    if value is None:
        return ""
    return str(value).strip()


def _as_purl(value: str) -> str | None:
    if value.lower().startswith("pkg:"):
        return value
    return None


def _identity(row: dict) -> tuple:
    if row["purl"]:
        return ("purl", row["purl"])
    return ("name", normalize_name(row["name"]), row["version"])


def _lookup_rows(rows: list[dict], http, matcher) -> tuple[dict[int, ComponentResult], list[str]]:
    cache: dict[tuple, ComponentResult] = {}
    resolved: dict[int, ComponentResult] = {}
    unresolved: list[str] = []
    for row in rows:
        key = _identity(row)
        if key not in cache:
            label = row["purl"] or f"{row['name']} {row['version']}".strip()
            print(f"lookup {label}", file=sys.stderr)
            cache[key] = _one(http, matcher, row)
        result = cache[key]
        resolved[row["row"]] = result
        if result.error:
            label = row["purl"] or f"{row['name']} {row['version']}".strip() or f"row {row['row']}"
            if label not in unresolved:
                unresolved.append(f"{label}: {result.error}")
    return resolved, unresolved


def _one(http, matcher, row: dict) -> ComponentResult:
    if not row["name"] and not row["purl"]:
        result = ComponentResult(name="", version=row["version"] or None, purl=None, vendor=None)
        result.error = "no component name or purl"
        return result
    try:
        return matcher(
            http,
            name=row["name"] or None,
            version=row["version"] or None,
            purl=row["purl"],
        )
    except CacheMiss as exc:
        result = ComponentResult(
            name=row["name"],
            version=row["version"] or None,
            purl=row["purl"],
            vendor=None,
        )
        result.error = f"offline cache miss ({exc.key})"
        return result
    except Exception as exc:  # one bad component should not drop the workbook
        result = ComponentResult(
            name=row["name"],
            version=row["version"] or None,
            purl=row["purl"],
            vendor=None,
        )
        result.error = str(exc)
        return result


def _write_result(worksheet, row: int, start: int, result: ComponentResult) -> None:
    notable = result.notable
    ids = " ".join(item.euvd_id for item in notable)
    aliases = []
    for item in notable:
        for alias in item.aliases:
            if alias.startswith(("CVE-", "GHSA-")) and alias not in aliases:
                aliases.append(alias)
    scores = [item.score for item in notable if item.score is not None]
    values = [
        ids,
        " ".join(aliases),
        format_score(max(scores) if scores else None),
        result.verdict,
        rollup_basis(result.records, result.version),
        _notes(result),
    ]
    for offset, value in enumerate(values):
        worksheet.cell(row, start + offset, value)


def _notes(result: ComponentResult) -> str:
    if result.error:
        return f"{result.error}; not marked affected"
    disagreements = [item for item in result.records if item.verdict == "disagree"]
    if result.verdict == "affected":
        text = f"{len(result.notable)} EUVD record(s) with a version basis. See EUVD Detail."
        if disagreements:
            text += f" {len(disagreements)} record(s) disagree between OSV and EUVD."
        return text
    if result.verdict == "disagree":
        return f"{len(disagreements)} record(s) disagree between OSV and EUVD. See EUVD Detail."
    if result.verdict == "not_affected":
        return f"{len(result.records)} EUVD record(s) do not include this version."
    if result.records and result.records[0].match_basis == "name-only":
        return "No version basis; not marked affected."
    if result.records:
        return result.records[0].evidence
    return "No EUVD record matched this component; not marked affected."


def _write_detail(workbook, resolved: dict[int, ComponentResult]) -> None:
    if "EUVD Detail" in workbook.sheetnames:
        del workbook["EUVD Detail"]
    sheet = workbook.create_sheet("EUVD Detail")
    sheet.append(list(DETAIL_HEADERS))
    link_font = Font(color="0563C1", underline="single")
    for row_number, result in resolved.items():
        records = result.records or [None]
        for item in records:
            if item is None:
                sheet.append(
                    [
                        result.name,
                        result.version or "",
                        result.purl or "",
                        "",
                        "",
                        "",
                        "",
                        result.verdict,
                        "name-only" if result.verdict == "unknown" else "",
                        "",
                        result.error or "No EUVD record matched this component.",
                        row_number,
                    ]
                )
                continue
            sheet.append(
                [
                    result.name,
                    result.version or "",
                    result.purl or "",
                    item.euvd_id,
                    " ".join(alias for alias in item.aliases if alias.startswith(("CVE-", "GHSA-", "PYSEC-"))),
                    item.severity or "",
                    item.score if item.score is not None else "",
                    item.verdict,
                    item.match_basis,
                    item.cross_check,
                    item.evidence,
                    row_number,
                ]
            )
            link = sheet.cell(sheet.max_row, 4)
            link.hyperlink = EUVD_LINK + item.euvd_id
            link.font = link_font
    sheet.column_dimensions["K"].width = 80
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{sheet.max_row}"
    sheet.freeze_panes = "A2"


def _write_run(workbook, resolved, unresolved, source: Path, http) -> None:
    if "EUVD Run" in workbook.sheetnames:
        del workbook["EUVD Run"]
    sheet = workbook.create_sheet("EUVD Run")
    counts = Counter(result.verdict for result in resolved.values())
    record_counts = Counter(item.verdict for result in resolved.values() for item in result.records)
    offline = bool(getattr(http, "offline", False))
    cache_dir = str(getattr(http, "cache_dir", ""))
    rows = [
        ("timestamp", datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")),
        ("tool_version", __version__),
        ("input", source.name),
        ("offline", str(offline).lower()),
        ("cache_dir", cache_dir),
        ("api_osv", "https://api.osv.dev/v1/query"),
        ("api_euvd_search", "https://euvdservices.enisa.europa.eu/api/search"),
        ("api_euvd_id", "https://euvdservices.enisa.europa.eu/api/enisaid"),
        ("components", str(len(resolved))),
        ("affected", str(counts["affected"])),
        ("not_affected", str(counts["not_affected"])),
        ("unknown", str(counts["unknown"])),
        ("disagree", str(counts["disagree"])),
        ("records_affected", str(record_counts["affected"])),
        ("records_not_affected", str(record_counts["not_affected"])),
        ("records_unknown", str(record_counts["unknown"])),
        ("records_disagree", str(record_counts["disagree"])),
    ]
    sheet.append(["key", "value"])
    for key, value in rows:
        sheet.append([key, value])
    sheet.append([])
    sheet.append(["components_not_resolved"])
    if unresolved:
        for item in unresolved:
            sheet.append([item])
    else:
        sheet.append(["none"])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 72


def _summary(results: list[ComponentResult], unresolved: list[str], target: Path) -> list[str]:
    counts = Counter(result.verdict for result in results)
    unique = []
    seen_components: set[tuple] = set()
    for result in results:
        marker = (result.purl or "", result.name, result.version or "")
        if marker in seen_components:
            continue
        seen_components.add(marker)
        unique.append(result)
    record_counts = Counter(item.verdict for result in unique for item in result.records)
    lines = [
        (
            f"components {len(results)} affected {counts['affected']} "
            f"not_affected {counts['not_affected']} unknown {counts['unknown']} "
            f"disagree {counts['disagree']}"
        ),
        (
            f"records affected {record_counts['affected']} "
            f"not_affected {record_counts['not_affected']} unknown {record_counts['unknown']} "
            f"disagree {record_counts['disagree']}"
        ),
    ]
    seen = set()
    for result in results:
        marker = (result.purl or "", result.name, result.version or "")
        if marker in seen:
            continue
        seen.add(marker)
        ids = " ".join(item.euvd_id for item in result.notable)
        lines.append(
            "\t".join(
                [
                    result.name,
                    result.version or "",
                    result.verdict,
                    rollup_basis(result.records, result.version),
                    ids,
                ]
            )
        )
    if unresolved:
        lines.append("unresolved " + str(len(unresolved)))
    lines.append(f"wrote {target}")
    return lines
