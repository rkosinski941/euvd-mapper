"""Insert worksheets into a copy of an xlsx without rewriting existing parts.

Black Duck reports carry charts, defined names, and autofilters. Opening the
file in a spreadsheet library and saving it rewrites those parts. This module
copies the zip and edits only the workbook index, relationships, and content
types, then appends the new sheet XML.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl.utils import get_column_letter

WORKSHEET_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
WORKSHEET_CONTENT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"

EUVD_WIDTHS = (12.73, 42.58, 11.79, 20.47, 20.47, 20.47, 30, 14, 12, 44, 10, 22, 22, 60.84, 20.47, 40, 40, 45)
SUMMARY_WIDTHS = (70, 90)


@dataclass
class SheetStyles:
    banner: int | None = None
    banner_rest: int | None = None
    header: int | None = None
    body: int | None = None
    component: int | None = None
    link: int | None = None
    wrap: int | None = None


def insert_euvd_sheets(
    source: Path,
    target: Path,
    banner: str,
    euvd_headers: list[str],
    euvd_rows: list[list],
    summary_rows: list[list],
    euvd_sheet: str = "EUVD Vulnerabilities",
    summary_sheet: str = "EUVD Run Summary",
    after_sheet: str = "Vulnerabilities",
) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as zin:
        workbook = zin.read("xl/workbook.xml").decode("utf-8")
        rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        styles = _styles_from_sheet(zin, workbook, rels, after_sheet)
        euvd_xml = _worksheet(
            [ [banner] + [""] * (len(euvd_headers) - 1), list(euvd_headers), *[_link_row(row) for row in euvd_rows] ],
            EUVD_WIDTHS,
            lambda row_number, column: _euvd_style(row_number, column, styles),
            freeze=True,
            autofilter=f"A2:{get_column_letter(len(euvd_headers))}2",
        )
        summary_xml = _worksheet(
            summary_rows,
            SUMMARY_WIDTHS,
            lambda row_number, column: _summary_style(row_number, column, summary_rows, styles),
            freeze=False,
            autofilter=None,
        )
        sheet_numbers = [
            int(match)
            for name in zin.namelist()
            for match in re.findall(r"^xl/worksheets/sheet(\d+)\.xml$", name)
        ]
        first = (max(sheet_numbers) if sheet_numbers else 0) + 1
        second = first + 1
        updated_workbook, _position = _insert_workbook(
            workbook,
            after_sheet,
            (
                (euvd_sheet, "rIdEUVD1", f"'{_defined_name(euvd_sheet)}'!$A$2:${get_column_letter(len(euvd_headers))}$2"),
                (summary_sheet, "rIdEUVD2", None),
            ),
        )
        updated_rels = _insert_rels(rels, (( "rIdEUVD1", first), ("rIdEUVD2", second)))
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
        updated_types = _insert_content_types(content_types, (first, second))
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    data = updated_workbook.encode("utf-8")
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    data = updated_rels.encode("utf-8")
                elif item.filename == "[Content_Types].xml":
                    data = updated_types.encode("utf-8")
                zout.writestr(item, data)
            zout.writestr(f"xl/worksheets/sheet{first}.xml", euvd_xml)
            zout.writestr(f"xl/worksheets/sheet{second}.xml", summary_xml)


def _link_row(row: list) -> list:
    copied = list(row)
    url = copied[14] if len(copied) > 14 else ""
    label = copied[3] if len(copied) > 3 else ""
    if isinstance(url, str) and url.startswith("http") and label:
        safe_url = url.replace('"', '""')
        safe_label = str(label).replace('"', '""')
        copied[14] = f'=HYPERLINK("{safe_url}","{safe_label}")'
    return copied


def _euvd_style(row_number: int, column: int, styles: SheetStyles) -> int | None:
    if row_number == 1:
        return styles.banner if column == 1 else styles.banner_rest
    if row_number == 2:
        return styles.header
    if column in {4, 15}:
        return styles.link if styles.link is not None else styles.body
    if column == 2:
        return styles.component if styles.component is not None else styles.body
    if column in {14, 16, 17, 18}:
        return styles.wrap if styles.wrap is not None else styles.body
    return styles.body


def _summary_style(row_number: int, column: int, rows: list[list], styles: SheetStyles) -> int | None:
    if row_number == 1:
        return styles.banner if column == 1 else styles.banner_rest
    label = ""
    if 1 <= row_number <= len(rows) and rows[row_number - 1]:
        label = str(rows[row_number - 1][0] or "")
    if label in {"Item", "Component"} or label.startswith("Excluded EUVD"):
        return styles.header
    return styles.component if styles.component is not None else styles.body


def _styles_from_sheet(zin: zipfile.ZipFile, workbook: str, rels: str, sheet_name: str) -> SheetStyles:
    path = _sheet_target(workbook, rels, sheet_name)
    if path is None or path not in zin.namelist():
        return SheetStyles()
    xml = zin.read(path).decode("utf-8")
    return SheetStyles(
        banner=_style_at(xml, "A1"),
        banner_rest=_style_at(xml, "B1") or _style_at(xml, "A1"),
        header=_style_at(xml, "A2"),
        body=_style_at(xml, "A3"),
        component=_style_at(xml, "B3") or _style_at(xml, "A3"),
        link=_style_at(xml, "D3") or _style_at(xml, "A3"),
        wrap=_style_at(xml, "J3") or _style_at(xml, "A3"),
    )


def _style_at(xml: str, ref: str) -> int | None:
    match = re.search(rf'<c r="{ref}"[^>]*\bs="(\d+)"', xml)
    if not match:
        return None
    return int(match.group(1))


def _sheet_target(workbook: str, rels: str, sheet_name: str) -> str | None:
    pattern = rf'<sheet\b[^>]*\bname="{re.escape(sheet_name)}"([^>]*)/?>'
    match = re.search(pattern, workbook)
    if not match:
        return None
    rid = re.search(r'\br:id="([^"]+)"', match.group(0))
    if not rid:
        return None
    relationship = re.search(rf'<Relationship\b[^>]*\bId="{re.escape(rid.group(1))}"[^>]*>', rels)
    if not relationship:
        return None
    target = re.search(r'\bTarget="([^"]+)"', relationship.group(0))
    if not target:
        return None
    path = target.group(1)
    if path.startswith("/"):
        return path.lstrip("/")
    if path.startswith("xl/"):
        return path
    return "xl/" + path


def _insert_workbook(xml: str, after_sheet: str, sheets: tuple[tuple[str, str, str | None], ...]) -> tuple[str, int]:
    names = re.findall(r'<sheet\b[^>]*\bname="([^"]+)"', xml)
    if after_sheet not in names:
        raise ValueError(f"sheet not found: {after_sheet}")
    position = names.index(after_sheet) + 1
    sheet_ids = [int(item) for item in re.findall(r'\bsheetId="(\d+)"', xml)]
    next_id = (max(sheet_ids) if sheet_ids else 0) + 1
    addition = []
    defined = []
    for offset, (name, rid, filter_ref) in enumerate(sheets):
        addition.append(f'<sheet name="{_xml_attr(name)}" sheetId="{next_id + offset}" r:id="{rid}"/>')
        if filter_ref:
            defined.append(
                f'<definedName name="_xlnm._FilterDatabase" localSheetId="{position + offset}" hidden="1">{filter_ref}</definedName>'
            )
    pattern = rf'(<sheet\b[^>]*\bname="{re.escape(after_sheet)}"[^>]*/>)'
    updated, count = re.subn(pattern, r"\1" + "".join(addition), xml, count=1)
    if count != 1:
        pattern = rf'(<sheet\b[^>]*\bname="{re.escape(after_sheet)}"[^>]*>\s*</sheet>)'
        updated, count = re.subn(pattern, r"\1" + "".join(addition), xml, count=1)
    if count != 1:
        raise ValueError(f"could not insert sheets after {after_sheet}")

    def _shift(match: re.Match) -> str:
        current = int(match.group(1))
        shifted = current + len(sheets) if current >= position else current
        return f'localSheetId="{shifted}"'

    updated = re.sub(r'localSheetId="(\d+)"', _shift, updated)
    for attr in ("activeTab", "firstSheet"):
        updated = re.sub(
            attr + r'="(\d+)"',
            lambda match, attr=attr: (
                f'{attr}="{int(match.group(1)) + len(sheets)}"'
                if int(match.group(1)) >= position
                else match.group(0)
            ),
            updated,
        )
    if defined:
        block = "".join(defined)
        if "</definedNames>" in updated:
            updated = updated.replace("</definedNames>", block + "</definedNames>", 1)
        else:
            updated = updated.replace("</workbook>", f"<definedNames>{block}</definedNames></workbook>", 1)
    return updated, position


def _insert_rels(xml: str, sheets: tuple[tuple[str, int], ...]) -> str:
    tags = []
    for rid, number in sheets:
        tags.append(f'<Relationship Id="{rid}" Target="worksheets/sheet{number}.xml" Type="{WORKSHEET_TYPE}"/>')
    return xml.replace("</Relationships>", "".join(tags) + "</Relationships>", 1)


def _insert_content_types(xml: str, numbers: tuple[int, ...]) -> str:
    tags = [
        f'<Override PartName="/xl/worksheets/sheet{number}.xml" ContentType="{WORKSHEET_CONTENT}"/>'
        for number in numbers
    ]
    return xml.replace("</Types>", "".join(tags) + "</Types>", 1)


def _worksheet(rows: list[list], widths: tuple[float, ...], style_for, freeze: bool, autofilter: str | None) -> str:
    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheetViews><sheetView workbookViewId="0">'
    ]
    if freeze:
        parts.append(
            '<pane ySplit="2" topLeftCell="A3" activePane="bottomLeft" state="frozen"/>'
            '<selection pane="bottomLeft" activeCell="A3" sqref="A3"/>'
        )
    parts.append('</sheetView></sheetViews><sheetFormatPr defaultRowHeight="12.9"/><cols>')
    parts.extend(
        f'<col min="{index}" max="{index}" width="{width}" customWidth="1"/>'
        for index, width in enumerate(widths, 1)
    )
    parts.append("</cols><sheetData>")
    for row_number, row in enumerate(rows, 1):
        height = ' ht="21.9" customHeight="1"' if row_number == 1 else ""
        cells = []
        for column, value in enumerate(row, 1):
            cells.append(_cell(f"{get_column_letter(column)}{row_number}", value, style_for(row_number, column)))
        parts.append(f'<row r="{row_number}"{height}>{"".join(cells)}</row>')
    parts.append("</sheetData>")
    if autofilter:
        parts.append(f'<autoFilter ref="{autofilter}"/>')
    parts.append(
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        "</worksheet>"
    )
    return "".join(parts)


def _cell(ref: str, value, style: int | None) -> str:
    style_attr = f' s="{style}"' if style is not None else ""
    if value is None or value == "":
        return f'<c r="{ref}"{style_attr}/>'
    if isinstance(value, bool):
        value = int(value)
    if isinstance(value, int) and not isinstance(value, bool):
        return f'<c r="{ref}"{style_attr}><v>{value}</v></c>'
    if isinstance(value, float):
        return f'<c r="{ref}"{style_attr}><v>{_number(value)}</v></c>'
    text = str(value)
    if text.startswith("=HYPERLINK"):
        formula = escape(text[1:])
        label = text.split('"')[3] if text.count('"') >= 4 else ""
        return f'<c r="{ref}"{style_attr} t="str"><f>{formula}</f><v>{escape(label)}</v></c>'
    return (
        f'<c r="{ref}"{style_attr} t="inlineStr"><is><t xml:space="preserve">{_xml_text(text)}</t></is></c>'
    )


def _number(value: float) -> str:
    return format(value, ".12g")


def _xml_text(value: str) -> str:
    cleaned = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)
    return escape(cleaned)


def _xml_attr(value: str) -> str:
    return escape(value, {'"': "&quot;"})


def _defined_name(sheet: str) -> str:
    return sheet.replace("'", "''")
