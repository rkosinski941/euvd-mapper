"""Build a synthetic Black Duck project report and check the enrich result."""

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from euvd_mapper.bdreport import BOM_HEADERS, VULN_HEADERS


def write_headers(sheet, titles):
    for col, value in enumerate(titles, 1):
        cell = sheet.cell(2, col, value)
        cell.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF5A2A82")
        cell.alignment = Alignment(horizontal="center")
    sheet.cell(1, 1, "Banner row for the report.")
    sheet.auto_filter.ref = "A2:%s2" % get_column_letter(len(titles))
    sheet.freeze_panes = "A3"


def build(path):
    book = Workbook()
    cover = book.active
    cover.title = "Summary"
    cover["A1"] = "Leave this sheet alone"
    cover["B2"] = "active tab stays here"
    vuln = book.create_sheet("Vulnerabilities")
    write_headers(vuln, VULN_HEADERS)
    vuln.cell(3, 1, "Critical")
    vuln.cell(3, 2, "SampleLib")
    vuln.cell(3, 3, "1.2.3")
    vuln.cell(3, 4, "CVE-2024-1000")
    notes = book.create_sheet("Notes")
    notes["A1"] = "keep me"
    bom = book.create_sheet("Bill of Materials")
    write_headers(bom, BOM_HEADERS)
    bom.cell(3, 1, "SampleLib")
    bom.cell(3, 2, "1.2.3")
    bom.cell(3, 3, "MIT")
    bom.cell(3, 4, "Managed Dependency - Binary")
    book.active = cover
    book.save(path)


def sheet_parts(path):
    found = {}
    with zipfile.ZipFile(path) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
        rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        pattern = re.compile(r'<sheet\b[^>]*\bname="([^"]+)"[^>]*>')
        for match in pattern.finditer(workbook):
            tag = match.group(0)
            rid = re.search(r'\br:id="([^"]+)"', tag)
            rel = re.search(
                r'<Relationship\b[^>]*\bId="%s"[^>]*>' % re.escape(rid.group(1)),
                rels,
            )
            target = re.search(r'\bTarget="([^"]+)"', rel.group(0)).group(1)
            if target.startswith("/"):
                target = target.lstrip("/")
            elif not target.startswith("xl/"):
                target = "xl/" + target
            found[match.group(1)] = archive.read(target)
    return found


def sheet_names(path):
    with zipfile.ZipFile(path) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
    return re.findall(r'<sheet\b[^>]*\bname="([^"]+)"', workbook)


def check(source, output):
    original = sheet_parts(source)
    updated = sheet_parts(output)
    for name, payload in original.items():
        if updated.get(name) != payload:
            raise SystemExit("original sheet changed: " + name)
    names = sheet_names(output)
    index = names.index("Vulnerabilities")
    if names[index + 1] != "EUVD Vulnerabilities":
        raise SystemExit("sheet order: " + " | ".join(names))
    if names[index + 2] != "EUVD Run Summary":
        raise SystemExit("sheet order: " + " | ".join(names))
    saved = load_workbook(output)
    if saved["Summary"]["A1"].value != "Leave this sheet alone":
        raise SystemExit("Summary A1 changed")
    if saved["Summary"]["B2"].value != "active tab stays here":
        raise SystemExit("Summary B2 changed")
    if saved["Notes"]["A1"].value != "keep me":
        raise SystemExit("Notes changed")
    if saved["Vulnerabilities"]["D3"].value != "CVE-2024-1000":
        raise SystemExit("Vulnerabilities changed")
    if saved["Bill of Materials"]["A3"].value != "SampleLib":
        raise SystemExit("Bill of Materials changed")
    print("sheets " + " | ".join(names))
    return 0


def main():
    mode = sys.argv[1]
    if mode == "build":
        build(Path(sys.argv[2]))
        print("built " + sys.argv[2])
        return 0
    if mode == "check":
        return check(Path(sys.argv[2]), Path(sys.argv[3]))
    raise SystemExit("usage: check_workbook.py build|check ...")


if __name__ == "__main__":
    raise SystemExit(main())
