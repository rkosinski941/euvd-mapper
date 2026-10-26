"""Black Duck project-report layout. These tests do not call the EUVD API."""

import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from euvd_mapper.bdreport import (
    BOM_HEADERS,
    EUVD_HEADERS,
    VULN_HEADERS,
    BomComponent,
    VulnRow,
    build_report,
    is_blackduck_project_report,
    severity_from_score,
    version_verdict,
)
from euvd_mapper.enrich import WorkbookError, enrich_workbook
from euvd_mapper.http import CacheMiss
from euvd_mapper.match import ComponentResult

ROOT = Path(__file__).resolve().parents[1]
WHEN = datetime(2026, 10, 26, 9, 0, tzinfo=ZoneInfo("America/New_York"))


class MapHttp:
    def __init__(self, routes):
        self.routes = routes
        self.offline = True
        self.cache_dir = "unused"
        self.calls = []

    def get_json(self, url):
        self.calls.append(url)
        if url not in self.routes:
            raise CacheMiss(url)
        return self.routes[url]


def _headers(sheet, headers):
    for col, value in enumerate(headers, 1):
        cell = sheet.cell(2, col, value)
        cell.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF5A2A82")
        cell.alignment = Alignment(horizontal="center")
    sheet.cell(1, 1, "Banner row for the report.").fill = PatternFill("solid", fgColor="FF9A73B3")
    sheet.cell(1, 2).fill = PatternFill("solid", fgColor="FF9A73B3")
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"
    sheet.freeze_panes = "A3"


def _project_workbook(path: Path) -> None:
    book = Workbook()
    cover = book.active
    cover.title = "Summary"
    cover["A1"] = "Leave this sheet alone"
    cover["B2"] = "active tab stays here"
    vuln = book.create_sheet("Vulnerabilities")
    _headers(vuln, VULN_HEADERS)
    vuln.cell(3, 1, "Critical")
    vuln.cell(3, 2, "SampleLib")
    vuln.cell(3, 3, "1.2.3")
    vuln.cell(3, 4, "CVE-2024-1000")
    vuln.cell(3, 5, "BDSA-1")
    vuln.cell(3, 10, "Wrapped description from the report.")
    after = book.create_sheet("Notes")
    _headers(after, ("Name", "Value"))
    after.cell(3, 1, "keep")
    after.cell(3, 2, "me")
    bom = book.create_sheet("Bill of Materials")
    _headers(bom, BOM_HEADERS)
    bom_rows = [
        ("Project Sample", "Unspecified", "Original Code"),
        ("SampleLib", "1.2.3", "Managed Dependency - Binary"),
        ("Widget :: Core", "2.0.0", "Transitive Dependency - Binary"),
    ]
    for row, (name, version, usage) in enumerate(bom_rows, 3):
        bom.cell(row, 1, name)
        bom.cell(row, 2, version)
        bom.cell(row, 3, "MIT")
        bom.cell(row, 4, usage)
    book.active = cover
    book.save(path)


def _routes():
    cve = {
        "id": "EUVD-2024-1000",
        "baseScore": 9.8,
        "baseScoreVersion": "3.1",
        "baseScoreVector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
        "epss": 0.22,
        "datePublished": "Oct 9, 2026, 10:18:56 AM",
        "dateUpdated": "Oct 15, 2026, 4:28:54 PM",
        "aliases": "CVE-2024-1000\nGHSA-aaaa-bbbb-cccc\n",
        "description": "SampleLib before 1.2.4 mishandles <tokens> & quotes.",
        "enisaIdProduct": [
            {"product": {"name": "SampleLib"}, "product_version": "1.0.0 <1.2.4"},
            {"product": {"name": "SampleLib"}, "product_version": "0 <1.0.0"},
        ],
    }
    out_of_range = {
        "id": "EUVD-2024-2000",
        "baseScore": 7.5,
        "baseScoreVersion": "3.1",
        "baseScoreVector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:H",
        "epss": 0.1,
        "aliases": "CVE-2024-2000",
        "description": "Old SampleLib only.",
        "enisaIdProduct": [
            {"product": {"name": "SampleLib"}, "product_version": "0 <1.0.0"},
        ],
    }
    malformed = {
        "id": "EUVD-2024-2001",
        "baseScore": 9.1,
        "aliases": "CVE-2024-2001",
        "description": "Broken EUVD range.",
        "enisaIdProduct": [
            {"product": {"name": "SampleLib"}, "product_version": "9.9.9, < 9.9.10"},
        ],
    }
    widget = {
        "id": "EUVD-2024-3000",
        "baseScore": 8.3,
        "baseScoreVersion": "4.0",
        "baseScoreVector": "CVSS:4.0/AV:N/AC:L/AT:N/PR:N/UI:N/VC:H/VI:H/VA:N/SC:N/SI:N/SA:N",
        "epss": 0.05,
        "datePublished": "Oct 1, 2026, 1:00:00 AM",
        "dateUpdated": "Oct 2, 2026, 1:00:00 AM",
        "aliases": "GHSA-cccc-dddd-eeee\nCVE-2024-3000",
        "description": "Widget before 2.1.0.",
        "enisaIdProduct": [
            {"product": {"name": "Apache Widget"}, "product_version": "1.0.0 <2.1.0"},
        ],
    }
    return {
        "https://euvdservices.enisa.europa.eu/api/enisaid?id=CVE-2024-1000": cve,
        "https://euvdservices.enisa.europa.eu/api/search?product=SampleLib&size=100&page=0": {
            "items": [cve, out_of_range, malformed],
            "total": 3,
        },
        "https://euvdservices.enisa.europa.eu/api/search?product=Widget&size=100&page=0": {
            "items": [widget],
            "total": 1,
        },
    }


class VerdictTest(unittest.TestCase):
    def test_severity_bands_come_from_the_euvd_score(self):
        self.assertEqual(severity_from_score(9), "Critical")
        self.assertEqual(severity_from_score(8.9), "High")
        self.assertEqual(severity_from_score(7), "High")
        self.assertEqual(severity_from_score(6.9), "Medium")
        self.assertEqual(severity_from_score(4), "Medium")
        self.assertEqual(severity_from_score(3.9), "Low")
        self.assertEqual(severity_from_score(0), "Low")
        self.assertEqual(severity_from_score(None), "Not Yet Rated")

    def test_simple_ranges_and_unparsed_strings(self):
        self.assertEqual(
            version_verdict("1.2.3", ["1.0.0 <1.2.4"]),
            ("affected (in range 1.0.0 <1.2.4)", "affected"),
        )
        self.assertEqual(version_verdict("1.2.3", ["1.0.0 <1.2.3"])[1], "not_affected")
        self.assertEqual(version_verdict("4.0.9", ["4.0.9"])[0], "affected (in range 4.0.9)")
        self.assertEqual(version_verdict("2.0.2", ["2.0.0 <2.2.0"])[0], "affected (in range 2.0.0 <2.2.0)")
        self.assertEqual(version_verdict("2.0.2", ["0 \u22641.2.20"])[1], "not_affected")
        label, kind = version_verdict("4.13.2", ["5.12.0, < 5.13.2"])
        self.assertEqual(kind, "unknown")
        self.assertIn("unparsed", label)
        self.assertEqual(version_verdict("1.2.3", [])[1], "unknown")
        self.assertEqual(version_verdict("2.14.1", ["2.0-beta9 <log4j-core*"])[1], "unknown")

    def test_rows_keep_cve_aliases_and_in_range_product_hits_only(self):
        report = build_report(
            MapHttp(_routes()),
            [VulnRow("SampleLib", "1.2.3", "CVE-2024-1000", "BDSA-1")],
            [
                BomComponent("Project Sample", "Unspecified", "Original Code"),
                BomComponent("SampleLib", "1.2.3", "Managed Dependency - Binary"),
                BomComponent("Widget :: Core", "2.0.0", "Transitive Dependency - Binary"),
            ],
            now=WHEN,
        )
        self.assertEqual(len(report.euvd_rows), 2)
        cve, extra = report.euvd_rows
        self.assertEqual(cve[0], "Critical")
        self.assertEqual(cve[1], "SampleLib")
        self.assertEqual(cve[3], "EUVD-2024-1000")
        self.assertEqual(cve[4], "CVE-2024-1000")
        self.assertEqual(cve[5], "BDSA-1")
        self.assertEqual(cve[6], "CVE-2024-1000; GHSA-aaaa-bbbb-cccc")
        self.assertEqual(cve[7], 9.8)
        self.assertEqual(cve[15], "affected (in range 1.0.0 <1.2.4)")
        self.assertEqual(cve[16], "CVE alias (from Vulnerabilities tab)")
        self.assertEqual(cve[17], "Also listed in Black Duck Vulnerabilities tab")
        self.assertEqual(extra[0], "High")
        self.assertEqual(extra[1], "Widget :: Core")
        self.assertEqual(extra[3], "EUVD-2024-3000")
        self.assertEqual(extra[8], "4.0")
        self.assertIn('fuzzy name match to EUVD product "Apache Widget"', extra[16])
        self.assertEqual(extra[17], "Not in Black Duck Vulnerabilities tab; review")
        excluded = [row[0] for row in report.summary_rows if str(row[0]).startswith("EUVD-2024-200")]
        self.assertEqual(excluded, ["EUVD-2024-2000 (SampleLib)", "EUVD-2024-2001 (SampleLib)"])
        joined = "\n".join(" ".join(str(cell) for cell in row) for row in report.summary_rows)
        self.assertNotIn("clean", joined.casefold())
        self.assertIn("still needs review", joined)


class WorkbookTest(unittest.TestCase):
    def setUp(self):
        self.root = Path("/tmp/euvd-bd-report")
        self.root.mkdir(exist_ok=True)
        self.source = self.root / "report.xlsx"
        self.output = self.root / "report-euvd.xlsx"
        _project_workbook(self.source)
        self.before = self.source.read_bytes()
        self.original_parts = _worksheet_parts(self.source)

    def test_inserts_sheets_and_keeps_original_parts(self):
        lines = enrich_workbook(self.source, self.output, MapHttp(_routes()), now=WHEN)
        self.assertEqual(self.source.read_bytes(), self.before)
        self.assertTrue(any(line.startswith("euvd_rows 2 ") for line in lines))
        saved = load_workbook(self.output)
        names = saved.sheetnames
        index = names.index("Vulnerabilities")
        self.assertEqual(names[index + 1], "EUVD Vulnerabilities")
        self.assertEqual(names[index + 2], "EUVD Run Summary")
        self.assertEqual(saved["Summary"]["A1"].value, "Leave this sheet alone")
        self.assertEqual(saved["Summary"]["B2"].value, "active tab stays here")
        self.assertEqual(saved["Bill of Materials"]["A3"].value, "Project Sample")
        self.assertEqual(saved["Vulnerabilities"]["D3"].value, "CVE-2024-1000")
        self.assertEqual(saved.active.title, "Summary")

        euvd = saved["EUVD Vulnerabilities"]
        self.assertEqual(euvd.freeze_panes, "A3")
        self.assertEqual(euvd.auto_filter.ref, "A2:R2")
        self.assertEqual([euvd.cell(2, col).value for col in range(1, 19)], list(EUVD_HEADERS))
        self.assertIn("public EUVD API on Oct 26 2026", euvd["A1"].value)
        self.assertEqual(euvd["A3"].value, "Critical")
        self.assertEqual(euvd["D3"].value, "EUVD-2024-1000")
        self.assertEqual(euvd["H3"].value, 9.8)
        self.assertIn("SampleLib before 1.2.4", euvd["N3"].value)
        self.assertIn("<tokens>", euvd["N3"].value)
        self.assertTrue(str(euvd["O3"].value).startswith("=HYPERLINK("))
        self.assertIn("EUVD-2024-1000", euvd["O3"].value)
        self.assertEqual(euvd["A4"].value, "High")
        self.assertEqual(euvd["B4"].value, "Widget :: Core")
        self.assertEqual(euvd.max_row, 4)

        summary = saved["EUVD Run Summary"]
        values = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(1, 15)}
        self.assertEqual(values["Run date"], "Oct 26 2026 (America/New_York)")
        self.assertEqual(values["BOM components checked (excluding Original Code)"], 2)
        self.assertEqual(values["Vulnerabilities tab rows (CVEs)"], 1)
        self.assertEqual(values["CVEs mapped to an EUVD ID"], 1)
        self.assertEqual(values["Additional EUVD entries via product search (not in Vulnerabilities tab)"], 1)
        self.assertEqual(values["Total rows on EUVD Vulnerabilities tab"], 2)
        self.assertEqual(
            values["Product-search hits excluded (version outside EUVD ranges, or fuzzy name without in-range version)"],
            2,
        )
        text = " ".join(
            str(summary.cell(row, 1).value)
            for row in range(1, summary.max_row + 1)
            if summary.cell(row, 1).value
        )
        self.assertNotIn("Project Sample", text)
        self.assertIn("SampleLib", text)
        self.assertIn("Widget :: Core", text)

        copied = _worksheet_parts(self.output)
        for name, payload in self.original_parts.items():
            self.assertEqual(copied[name], payload, name)
        _assert_filter_shift(self.source, self.output)
        _assert_header_style_copied(self.output)

    def test_column_flags_keep_the_generic_layout(self):
        def matcher(http, name=None, version=None, purl=None, vendor=None):
            return ComponentResult(name=name or "", version=version, purl=purl, vendor=None)

        enrich_workbook(
            self.source,
            self.output,
            MapHttp({}),
            sheet="Bill of Materials",
            name_column="Component",
            version_column="Version",
            matcher=matcher,
            now=WHEN,
        )
        saved = load_workbook(self.output)
        self.assertNotIn("EUVD Vulnerabilities", saved.sheetnames)
        self.assertIn("EUVD Detail", saved.sheetnames)
        self.assertEqual(saved["Bill of Materials"].cell(2, 14).value, "EUVD IDs")

    def test_refuses_to_overwrite_the_input(self):
        with self.assertRaises(WorkbookError):
            enrich_workbook(self.source, self.source, MapHttp({}), now=WHEN)

    def test_single_space_bdsa_header_is_not_this_layout(self):
        path = self.root / "other.xlsx"
        book = Workbook()
        bom = book.active
        bom.title = "Bill of Materials"
        for col, value in enumerate(BOM_HEADERS, 1):
            bom.cell(2, col, value)
        vuln = book.create_sheet("Vulnerabilities")
        headers = list(VULN_HEADERS)
        headers[19] = "BDSA Overall (Temporal)"
        for col, value in enumerate(headers, 1):
            vuln.cell(2, col, value)
        book.save(path)
        self.assertFalse(is_blackduck_project_report(load_workbook(path)))

    def test_existing_samples_are_not_project_reports(self):
        self.assertFalse(is_blackduck_project_report(load_workbook(ROOT / "examples" / "black-duck-sample.xlsx")))
        self.assertFalse(is_blackduck_project_report(load_workbook(ROOT / "examples" / "cyclonedx-sample.xlsx")))

    def test_offline_miss_is_reported_and_not_called_a_miss_of_vulnerabilities(self):
        report = build_report(
            MapHttp({}),
            [VulnRow("SampleLib", "1.2.3", "CVE-2024-1000", "")],
            [BomComponent("SampleLib", "1.2.3", "Managed Dependency - Binary")],
            now=WHEN,
        )
        self.assertEqual(report.euvd_rows[0][3], "no EUVD match")
        self.assertIn("offline cache miss", report.euvd_rows[0][17])
        self.assertNotIn("clean", report.euvd_rows[0][17].casefold())


def _worksheet_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        }


def _assert_filter_shift(source: Path, output: Path) -> None:
    import re

    def filters(path):
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("xl/workbook.xml").decode("utf-8")
        found = {}
        for match in re.finditer(r'<definedName\b([^>]*)>([^<]*)</definedName>', xml):
            attrs = match.group(1)
            if "_xlnm._FilterDatabase" not in attrs:
                continue
            sheet_id = re.search(r'localSheetId="(\d+)"', attrs)
            if sheet_id:
                found[match.group(2)] = int(sheet_id.group(1))
        return found

    before = filters(source)
    after = filters(output)
    # Vulnerabilities is the second sheet (index 1). Sheets at index >= 2 shift by 2.
    for ref, sheet_id in before.items():
        if ref not in after:
            raise AssertionError(f"missing defined name {ref} in {after}")
        if "Vulnerabilities" in ref and "EUVD" not in ref:
            if after[ref] != sheet_id:
                raise AssertionError(f"{ref} localSheetId changed from {sheet_id} to {after[ref]}")
        elif sheet_id >= 2:
            if after[ref] != sheet_id + 2:
                raise AssertionError(f"{ref} localSheetId {sheet_id} did not shift to {sheet_id + 2} ({after})")
    euvd = [ref for ref in after if "EUVD Vulnerabilities" in ref]
    if len(euvd) != 1 or after[euvd[0]] != 2:
        raise AssertionError(after)


def _assert_header_style_copied(path: Path) -> None:
    import re

    with zipfile.ZipFile(path) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
        rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")

        def sheet_xml(name):
            tag = re.search(rf'<sheet\b[^>]*name="{re.escape(name)}"[^>]*>', workbook).group(0)
            rid = re.search(r'r:id="([^"]+)"', tag).group(1)
            rel = re.search(rf'<Relationship\b[^>]*\bId="{re.escape(rid)}"[^>]*>', rels).group(0)
            target = re.search(r'Target="([^"]+)"', rel).group(1).lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            return archive.read(target).decode("utf-8")

        vuln = sheet_xml("Vulnerabilities")
        euvd = sheet_xml("EUVD Vulnerabilities")
        for ref in ("A1", "A2", "A3", "B3", "D3"):
            left = re.search(rf'<c r="{ref}"[^>]*s="(\d+)"', vuln)
            right = re.search(rf'<c r="{ref}"[^>]*s="(\d+)"', euvd)
            if left and right and left.group(1) != right.group(1):
                raise AssertionError(f"{ref} style {left.group(1)} was not copied ({right.group(1)})")


class TrimTest(unittest.TestCase):
    def test_cache_keeps_workbook_fields(self):
        from euvd_mapper.http import trim_euvd_response

        trimmed = trim_euvd_response(
            {
                "id": "EUVD-2024-1000",
                "description": "text",
                "baseScore": 9.8,
                "baseScoreVersion": "3.1",
                "baseScoreVector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
                "epss": 0.22,
                "datePublished": "Oct 9, 2026, 10:18:56 AM",
                "dateUpdated": "Oct 15, 2026, 4:28:54 PM",
                "aliases": "CVE-2024-1000",
                "references": "drop me",
                "enisaIdProduct": [
                    {"product_version": "1.0.0 <1.2.4", "product": {"name": "SampleLib", "vendor": {"name": "Example"}}}
                ],
            }
        )
        self.assertEqual(trimmed["baseScoreVersion"], "3.1")
        self.assertEqual(trimmed["baseScoreVector"].startswith("CVSS:3.1/"), True)
        self.assertEqual(trimmed["datePublished"], "Oct 9, 2026, 10:18:56 AM")
        self.assertNotIn("references", trimmed)


if __name__ == "__main__":
    unittest.main()
