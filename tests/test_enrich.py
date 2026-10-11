import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from euvd_mapper.enrich import WorkbookError, enrich_workbook
from euvd_mapper.match import ComponentResult, RecordHit


class FakeHttp:
    offline = True
    cache_dir = "unused"


def _hit(eid, verdict="affected", basis="osv-range", cross="agree", score=9.8, aliases=None):
    return RecordHit(
        euvd_id=eid,
        aliases=aliases or ["CVE-2021-44228"],
        score=score,
        severity="critical" if score and score >= 9 else "high",
        verdict=verdict,
        match_basis=basis,
        evidence=f"evidence for {eid}",
        cross_check=cross,
    )


class EnrichTest(unittest.TestCase):
    def test_appends_columns_and_keeps_the_original_workbook(self):
        source = Path(self.id() + ".xlsx")
        # unique temp next to cwd would collide; use a directory under /tmp via Path
        root = Path("/tmp/euvd-enrich-test")
        root.mkdir(exist_ok=True)
        source = root / "bom.xlsx"
        output = root / "bom-euvd.xlsx"
        book = Workbook()
        cover = book.active
        cover.title = "Cover"
        cover["A1"] = "Public sample"
        cover["B2"] = "do not touch"
        sheet = book.create_sheet("BOM Components")
        sheet["A1"] = "Black Duck BOM component report"
        sheet["A2"] = "Project public-sample"
        headers = [
            "Component name",
            "Component version name",
            "Version",
            "License",
            "Origin name",
            "Origin id",
            "Package URL",
        ]
        for col, value in enumerate(headers, 1):
            sheet.cell(3, col, value)
        sheet.cell(4, 1, "log4j-core")
        sheet.cell(4, 2, "2.14.1")
        sheet.cell(4, 3, "2.14.1")
        sheet.cell(4, 7, "pkg:maven/org.apache.logging.log4j/log4j-core@2.14.1")
        sheet.cell(5, 1, "log4j-core")
        sheet.cell(5, 2, "2.14.1")
        sheet.cell(5, 7, "pkg:maven/org.apache.logging.log4j/log4j-core@2.14.1")
        sheet.cell(6, 1, "Apache Tomcat")
        sheet.cell(6, 2, "latest")
        sheet.cell(6, 3, "latest")
        book.save(source)
        before = source.read_bytes()

        calls = []

        def matcher(http, name=None, version=None, purl=None, vendor=None):
            calls.append((name, version, purl))
            if purl and "log4j" in purl:
                result = ComponentResult(name=name or "log4j-core", version="2.14.1", purl=purl, vendor=None)
                result.records = [_hit("EUVD-2021-34768")]
                return result
            result = ComponentResult(name=name or "", version=version, purl=purl, vendor=None)
            result.records = [
                RecordHit(
                    euvd_id="EUVD-2022-4051",
                    aliases=["CVE-2022-25762"],
                    score=None,
                    severity=None,
                    verdict="unknown",
                    match_basis="name-only",
                    evidence="version 'latest' is not a comparable version",
                    cross_check="not_run",
                )
            ]
            return result

        lines = enrich_workbook(source, output, FakeHttp(), matcher=matcher)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(len(calls), 2)
        text = "\n".join(lines)
        self.assertIn("components 3 affected 2 not_affected 0 unknown 1 disagree 0", text)
        self.assertIn("EUVD-2021-34768", text)

        saved = load_workbook(output)
        self.assertEqual(saved["Cover"]["A1"].value, "Public sample")
        self.assertEqual(saved["Cover"]["B2"].value, "do not touch")
        bom = saved["BOM Components"]
        self.assertEqual(bom["A1"].value, "Black Duck BOM component report")
        self.assertEqual(bom["A4"].value, "log4j-core")
        self.assertEqual(bom["G4"].value, "pkg:maven/org.apache.logging.log4j/log4j-core@2.14.1")
        self.assertEqual(bom.cell(3, 8).value, "EUVD IDs")
        self.assertEqual(bom.cell(3, 11).value, "Version verdict")
        self.assertEqual(bom.cell(3, 13).value, "Evidence/notes")
        self.assertEqual(bom.cell(4, 8).value, "EUVD-2021-34768")
        self.assertEqual(bom.cell(4, 11).value, "affected")
        self.assertEqual(bom.cell(4, 12).value, "osv-range")
        self.assertEqual(bom.cell(5, 8).value, "EUVD-2021-34768")
        self.assertEqual(bom.cell(6, 11).value, "unknown")
        self.assertIn(bom.cell(6, 8).value, (None, ""))
        detail = saved["EUVD Detail"]
        self.assertEqual(detail["A1"].value, "Component name")
        self.assertEqual(detail["D2"].value, "EUVD-2021-34768")
        self.assertIn("euvd.enisa.europa.eu/vulnerability/EUVD-2021-34768", detail["D2"].hyperlink.target)
        run = saved["EUVD Run"]
        values = {run.cell(row, 1).value: run.cell(row, 2).value for row in range(2, 20)}
        self.assertEqual(values["tool_version"], "0.2.0")
        self.assertEqual(values["affected"], "2")
        self.assertEqual(values["unknown"], "1")
        self.assertIn("api.osv.dev", values["api_osv"])

    def test_refuses_to_overwrite_the_input(self):
        root = Path("/tmp/euvd-enrich-test")
        root.mkdir(exist_ok=True)
        source = root / "same.xlsx"
        Workbook().save(source)
        with self.assertRaises(WorkbookError):
            enrich_workbook(source, source, FakeHttp(), matcher=lambda *args, **kwargs: None)


if __name__ == "__main__":
    unittest.main()
