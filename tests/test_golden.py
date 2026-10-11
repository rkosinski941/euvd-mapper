import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from euvd_mapper.enrich import enrich_workbook
from euvd_mapper.http import HttpCache
from euvd_mapper.versionrange import evaluate_product_version

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "tests" / "fixtures" / "cache"
STDOUT = ROOT / "tests" / "fixtures" / "demo-stdout.txt"


class SharedRangeTest(unittest.TestCase):
    def test_python_matches_shared_range_cases(self):
        cases = json.loads((ROOT / "tests" / "fixtures" / "ranges.json").read_text(encoding="utf-8"))
        for item in cases:
            got = evaluate_product_version(item["range"], item["version"]).verdict
            self.assertEqual(got, item["verdict"], item)


class GoldenBomTest(unittest.TestCase):
    def test_offline_black_duck_sample(self):
        output = Path("/tmp/euvd-golden-black-duck.xlsx")
        lines = enrich_workbook(
            ROOT / "examples" / "black-duck-sample.xlsx",
            output,
            HttpCache(CACHE, offline=True, min_interval=0),
        )
        expected = STDOUT.read_text(encoding="utf-8").splitlines()
        actual = [line for line in lines if not line.startswith("wrote ")]
        frozen = [line for line in expected if not line.startswith("wrote ")]
        self.assertEqual(actual, frozen)

        book = load_workbook(output)
        self.assertEqual(book["Cover"]["A1"].value, "Public sample")
        self.assertEqual(book["BOM Components"]["A4"].value, "tomcat-catalina")
        detail = {
            (row[0], row[1], row[3]): row
            for row in book["EUVD Detail"].iter_rows(min_row=2, values_only=True)
        }
        tomcat = detail[("tomcat-catalina", "9.0.20", "EUVD-2022-4051")]
        self.assertEqual(tomcat[7], "disagree")
        self.assertEqual(tomcat[8], "euvd-range")
        self.assertIn("includes 9.0.20", tomcat[10])
        self.assertIn("did not list", tomcat[10])

        fixed = detail[("Apache Tomcat", "9.0.21", "EUVD-2022-4051")]
        self.assertEqual(fixed[7], "not_affected")
        self.assertIn("excludes 9.0.21", fixed[10])

        latest = detail[("Apache Tomcat", "latest", "EUVD-2022-4051")]
        self.assertEqual(latest[7], "unknown")
        self.assertEqual(latest[8], "name-only")
        self.assertNotEqual(latest[7], "affected")

        log4shell = detail[("log4j-core", "2.14.1", "EUVD-2021-34768")]
        self.assertEqual(log4shell[7], "affected")
        self.assertEqual(log4shell[8], "osv-range")
        self.assertIn("not a parseable range", log4shell[10])
        self.assertIn("CVE-2021-44228", log4shell[4])

        lodash = detail[("lodash", "4.17.15", "EUVD-2021-0912")]
        self.assertEqual(lodash[7], "affected")
        self.assertEqual(lodash[9], "agree")
        self.assertIn("CVE-2021-23337", lodash[4])

        jackson = detail[("jackson-databind", "2.9.10", "EUVD-2020-0440")]
        self.assertEqual(jackson[7], "unknown")
        self.assertIn("contradicts the description", jackson[10])

        django = detail[("Django", "3.2.0", "EUVD-2021-0067")]
        self.assertEqual(django[7], "affected")
        self.assertEqual(django[8], "osv-range")
        self.assertIn("CVE-2021-33203", django[4])
        self.assertIn("no product_version", django[10])

        sheet = book["BOM Components"]
        self.assertEqual(sheet.cell(3, 8).value, "EUVD IDs")
        self.assertEqual(sheet.cell(4, 11).value, "affected")
        self.assertIn("EUVD-2022-4051", sheet.cell(4, 8).value)
        self.assertEqual(sheet.cell(10, 11).value, "unknown")

    def test_offline_cyclonedx_sample_keeps_log4shell(self):
        output = Path("/tmp/euvd-golden-cyclonedx.xlsx")
        enrich_workbook(
            ROOT / "examples" / "cyclonedx-sample.xlsx",
            output,
            HttpCache(CACHE, offline=True, min_interval=0),
        )
        book = load_workbook(output)
        detail = {
            (row[0], row[3]): row[7]
            for row in book["EUVD Detail"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(detail[("log4j-core", "EUVD-2021-34768")], "affected")
        self.assertEqual(detail[("django", "EUVD-2021-0061")], "affected")
        self.assertEqual(detail[("Apache Tomcat", "EUVD-2022-4051")], "not_affected")
